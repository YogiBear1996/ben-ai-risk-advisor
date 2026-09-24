"""Document loaders: PDF, DOCX, Markdown and XLSX -> titled sections with page numbers.

Sections follow the document's own structure (headings, articles, clauses) so that chunks can
be cited precisely.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt", ".xlsx"}

# Lines that start a new clause/article/section in regulatory and standards text.
CLAUSE_RE = re.compile(
    r"""^(
        (?:Article|Art\.|Clause|Section|Annex|Chapter|Part|Principle|Recital|Schedule)
            \s+[0-9IVXLC]+[A-Za-z]?(?:\.\d+)*\b.*
      | \d{1,2}(?:\.\d{1,2}){0,3}\.?\s+[A-Z][^.]{2,90}
    )$""",
    re.VERBOSE,
)
MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")


@dataclass
class Section:
    heading: str | None
    text: str
    page: int | None = None


@dataclass
class LoadedDocument:
    title: str
    sections: list[Section] = field(default_factory=list)


def humanise_filename(path: Path) -> str:
    return re.sub(r"[_\-]+", " ", path.stem).strip().title()


def is_clause_heading(line: str) -> bool:
    line = line.strip()
    return 3 <= len(line) <= 120 and bool(CLAUSE_RE.match(line))


def _split_lines_on_clauses(
    lines: list[str], page: int | None, current: Section | None, out: list[Section]
) -> Section | None:
    """Append lines to `current`, starting a new Section at each clause heading."""
    for line in lines:
        stripped = line.strip()
        if is_clause_heading(stripped):
            if current and current.text.strip():
                out.append(current)
            current = Section(heading=stripped, text="", page=page)
            continue
        if current is None:
            current = Section(heading=None, text="", page=page)
        if not current.text.strip() and current.page is None:
            current.page = page
        current.text += line + "\n"
    return current


def load_markdown(path: Path) -> LoadedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    title: str | None = None
    sections: list[Section] = []
    current = Section(heading=None, text="")
    for line in text.splitlines():
        m = MD_HEADING_RE.match(line)
        if m:
            heading = m.group(2).strip()
            if title is None and len(m.group(1)) == 1:
                title = heading
                continue
            if current.text.strip():
                sections.append(current)
            current = Section(heading=heading, text="")
        else:
            current.text += line + "\n"
    if current.text.strip():
        sections.append(current)
    return LoadedDocument(title=title or humanise_filename(path), sections=sections)


def load_docx(path: Path) -> LoadedDocument:
    import docx

    document = docx.Document(str(path))
    title = (document.core_properties.title or "").strip() or None
    sections: list[Section] = []
    current: Section | None = None
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style == "title":
            title = title or text
            continue
        if style.startswith("heading") or is_clause_heading(text):
            if title is None and style == "heading 1" and not sections and current is None:
                title = text
                continue
            if current and current.text.strip():
                sections.append(current)
            current = Section(heading=text, text="")
            continue
        if current is None:
            current = Section(heading=None, text="")
        current.text += text + "\n"
    if current and current.text.strip():
        sections.append(current)
    # Tables (common in policy docs) become their own sections.
    for i, table in enumerate(document.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        body = "\n".join(r for r in rows if r.strip(" |"))
        if body:
            sections.append(Section(heading=f"Table {i}", text=body))
    return LoadedDocument(title=title or humanise_filename(path), sections=sections)


def load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    meta_title = None
    if reader.metadata and reader.metadata.title:
        meta_title = str(reader.metadata.title).strip() or None
    sections: list[Section] = []
    current: Section | None = None
    for page_no, page in enumerate(reader.pages, start=1):
        lines = (page.extract_text() or "").splitlines()
        current = _split_lines_on_clauses(lines, page_no, current, sections)
    if current and current.text.strip():
        sections.append(current)
    return LoadedDocument(title=meta_title or humanise_filename(path), sections=sections)


def load_text(path: Path) -> LoadedDocument:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    sections: list[Section] = []
    current = _split_lines_on_clauses(lines, None, None, sections)
    if current and current.text.strip():
        sections.append(current)
    return LoadedDocument(title=humanise_filename(path), sections=sections)


def load_xlsx_as_text(path: Path) -> LoadedDocument:
    """Generic spreadsheet (not a control library): one section per row."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sections: list[Section] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        for i, row in enumerate(rows[1:], start=2):
            pairs = [
                f"{h or f'Column {j + 1}'}: {v}"
                for j, (h, v) in enumerate(zip(header, row, strict=False))
                if v not in (None, "")
            ]
            if pairs:
                sections.append(Section(heading=f"{ws.title} row {i}", text="\n".join(pairs)))
    wb.close()
    return LoadedDocument(title=humanise_filename(path), sections=sections)


def load_document(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return load_markdown(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".txt":
        return load_text(path)
    if suffix == ".xlsx":
        return load_xlsx_as_text(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")
