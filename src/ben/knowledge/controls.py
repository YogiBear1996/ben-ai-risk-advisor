"""AI Risk & Control Library: XLSX -> structured ControlRecords.

Column names are matched case/space-insensitively with common synonyms; extra columns are kept.
"""

from __future__ import annotations

import re
from pathlib import Path

from ben.knowledge.types import ControlRecord

CONTROL_LIBRARY_TITLE = "AI Risk & Control Library"

COLUMN_SYNONYMS: dict[str, set[str]] = {
    "control_id": {"control_id", "controlid", "control_ref", "control_reference", "control_no"},
    "risk_category": {"risk_category", "category", "risk_domain", "risk_area"},
    "ai_type": {"ai_type", "ai_category", "type_of_ai", "aitype"},
    "risk_description": {"risk_description", "risk", "risk_statement"},
    "control_description": {"control_description", "control", "control_statement"},
    "owner": {"owner", "control_owner", "responsible"},
    "framework_mapping": {"framework_mapping", "frameworks", "mapping", "framework"},
}


def _norm(header: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")


def normalise_ai_type(value: object) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    v = str(value).strip().lower().replace("-", "").replace(" ", "")
    if v in {"both", "all", "traditional/genai", "traditional&genai", "genai/traditional"}:
        return "Both"
    if "gen" in v or "llm" in v:
        return "GenAI"
    if "trad" in v or v in {"ml", "predictive"}:
        return "Traditional"
    return str(value).strip()


def _map_header(header_row: tuple) -> dict[int, str] | None:
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(header_row):
        key = _norm(raw)
        if not key:
            continue
        field = next((f for f, syn in COLUMN_SYNONYMS.items() if key in syn), None)
        mapping[idx] = field if field and field not in mapping.values() else f"extra:{raw}"
    return mapping if "control_id" in mapping.values() else None


def load_controls(path: Path) -> list[ControlRecord]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    records: list[ControlRecord] = []
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header_idx, mapping = None, None
            for i, row in enumerate(rows[:10]):
                mapping = _map_header(row)
                if mapping:
                    header_idx = i
                    break
            if header_idx is None or mapping is None:
                continue
            for row in rows[header_idx + 1 :]:
                values: dict[str, str] = {}
                extra: dict[str, str] = {}
                for idx, field in mapping.items():
                    if idx >= len(row) or row[idx] in (None, ""):
                        continue
                    value = str(row[idx]).strip()
                    if field.startswith("extra:"):
                        extra[field.removeprefix("extra:")] = value
                    else:
                        values[field] = value
                if not values.get("control_id"):
                    continue
                records.append(
                    ControlRecord(
                        control_id=values["control_id"],
                        risk_category=values.get("risk_category"),
                        ai_type=normalise_ai_type(values.get("ai_type")),
                        risk_description=values.get("risk_description"),
                        control_description=values.get("control_description"),
                        owner=values.get("owner"),
                        framework_mapping=values.get("framework_mapping"),
                        extra=extra,
                        source_path=path.name,
                    )
                )
    finally:
        wb.close()
    return records
