"""Generate the sample DOCX, PDF and XLSX files in ./knowledge.

Run with: uv run python scripts/make_sample_knowledge.py
All content is placeholder paraphrase for testing, not authoritative text.
"""

from __future__ import annotations

from pathlib import Path

import docx
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

KNOWLEDGE = Path(__file__).resolve().parents[1] / "knowledge"

ISO_CLAUSES = [
    (
        "Clause 4 Context of the organization",
        "The organization determines internal and external issues relevant to its purpose "
        "and the intended purpose of the AI systems it develops, provides or uses. It "
        "determines its roles with respect to AI systems (for example AI provider, AI producer "
        "or AI customer), the needs of interested parties, and the scope of the AI management "
        "system (AIMS).",
    ),
    (
        "Clause 5 Leadership",
        "Top management demonstrates leadership and commitment by establishing an AI policy "
        "aligned with the organization's strategic direction, integrating AIMS requirements "
        "into business processes, and assigning roles, responsibilities and authorities.",
    ),
    (
        "Clause 6.1.2 AI risk assessment",
        "The organization defines and applies an AI risk assessment process that identifies "
        "risks that could prevent achieving its AI objectives, analyses the potential "
        "consequences to the organization, individuals and societies, assesses likelihood, "
        "and compares results against defined risk criteria to prioritise treatment.",
    ),
    (
        "Clause 6.1.3 AI risk treatment",
        "The organization selects appropriate AI risk treatment options and determines the "
        "controls necessary to implement them, compares them with the reference controls in "
        "Annex A, and produces a Statement of Applicability justifying inclusion or exclusion "
        "of controls. Residual risks are accepted by the risk owners.",
    ),
    (
        "Clause 6.1.4 AI system impact assessment",
        "The organization defines a process for assessing the potential consequences for "
        "individuals, groups of individuals and societies that can result from the "
        "development, provision or use of AI systems. The assessment considers the technical "
        "and societal context and the jurisdictions where the system is deployed. Results are "
        "documented and are considered in the risk assessment. Where appropriate, results are "
        "made available to relevant interested parties.",
    ),
    (
        "Clause 7 Support",
        "The organization provides the resources, competence, awareness, communication and "
        "documented information needed for the AIMS, including competence of people working "
        "on AI systems.",
    ),
    (
        "Clause 8 Operation",
        "The organization plans, implements and controls the processes needed to meet AIMS "
        "requirements. It performs AI risk assessments (8.2), implements the AI risk treatment "
        "plan (8.3) and performs AI system impact assessments (8.4) at planned intervals or "
        "when significant changes are proposed or occur, retaining documented results.",
    ),
    (
        "Clause 9 Performance evaluation",
        "The organization monitors, measures, analyses and evaluates the AIMS, conducts "
        "internal audits at planned intervals, and top management reviews the AIMS.",
    ),
    (
        "Clause 10 Improvement",
        "The organization continually improves the AIMS and reacts to nonconformities by "
        "taking corrective action.",
    ),
    (
        "Annex A Reference control objectives and controls",
        "Annex A lists reference controls grouped by objective, including: policies related to "
        "AI; internal organization; resources for AI systems; assessing impacts of AI systems; "
        "AI system life cycle; data for AI systems; information for interested parties; use of "
        "AI systems; and third-party and customer relationships.",
    ),
]

PDPL_PAGES = [
    [
        (
            "Section 1 Purpose and status",
            "These internal guidance notes summarise how the UAE "
            "Personal Data Protection Law (Federal Decree-Law No. 45 of 2021, PDPL) should be "
            "considered when AI systems process personal data. They are a placeholder for testing "
            "and do not quote the law. Article references must be confirmed by Legal.",
        ),
        (
            "Section 2 Scope",
            "The PDPL applies to the processing of personal data of data "
            "subjects residing or having a place of business in the UAE, and to controllers and "
            "processors in the UAE, whether or not processing takes place in the UAE. It does not "
            "apply to data regulated by certain sector-specific or free-zone data protection "
            "regimes (for example DIFC and ADGM have their own laws).",
        ),
    ],
    [
        (
            "Section 3 Lawful basis and consent",
            "Processing generally requires the data subject's "
            "consent, which must be specific, clear and unambiguous and capable of being withdrawn, "
            "unless an exception applies, such as processing necessary to perform a contract with "
            "the data subject or to comply with legal obligations. Training an AI model on "
            "passenger data is a new purpose; the team must confirm whether an existing basis "
            "covers it or whether consent or another exception is needed.",
        ),
        (
            "Section 4 Processing principles",
            "Personal data must be processed fairly, "
            "transparently and lawfully; collected for specific and clear purposes and not "
            "further processed incompatibly; adequate and limited to what is necessary; accurate "
            "and kept up to date; kept securely; and not retained longer than necessary. "
            "Anonymised data falls outside these obligations, so anonymisation before model "
            "training is strongly preferred.",
        ),
    ],
    [
        (
            "Section 5 Automated processing and data subject rights",
            "Data subjects have rights "
            "including access, correction, erasure, restriction and objection. The PDPL gives data "
            "subjects a right to object to decisions based solely on automated processing, "
            "including profiling, that have legal consequences or seriously affect them. AI use "
            "cases that make or materially influence decisions about passengers or staff need a "
            "human review route.",
        ),
        (
            "Section 6 Security, impact assessment and cross-border transfer",
            "Controllers must "
            "apply appropriate technical and organisational security measures, notify the UAE Data "
            "Office of breaches that affect privacy, and carry out a data protection impact "
            "assessment before processing likely to pose a high risk, such as large-scale "
            "processing of sensitive data or systematic evaluation using new technologies like AI. "
            "Transfers of personal data outside the UAE are permitted to countries with adequate "
            "protection or under specified safeguards and exceptions; model training on "
            "third-party cloud infrastructure abroad must be reviewed against these rules.",
        ),
    ],
]

CONTROLS = [
    (
        "AIRC-001",
        "Governance & Accountability",
        "Both",
        "AI systems deployed without a named accountable owner or approval.",
        "Every AI use case is registered in the AI inventory with a named business owner and "
        "approved by the AI governance forum before production use.",
        "Head of AI Governance",
        "ISO/IEC 42001 Cl.5; EU AI Act Art. 26",
    ),
    (
        "AIRC-002",
        "Risk Classification",
        "Both",
        "AI use cases are not classified against regulatory risk tiers, so obligations are missed.",
        "Each use case is classified against EU AI Act risk tiers (prohibited, high-risk, limited, "
        "minimal) and local regulation at intake, with the rationale documented.",
        "AI Governance Office",
        "EU AI Act Art. 5, Art. 6, Annex III",
    ),
    (
        "AIRC-003",
        "Impact Assessment",
        "Both",
        "Harms to passengers, staff or society are not identified before deployment.",
        "An AI system impact assessment is completed and approved before deployment and on "
        "significant change, covering affected groups, fundamental rights and mitigations.",
        "Use Case Owner",
        "ISO/IEC 42001 Cl.6.1.4, Cl.8.4; EU AI Act Art. 27",
    ),
    (
        "AIRC-004",
        "Data Governance & Privacy",
        "Both",
        "Personal data used to train or run models without a lawful basis or minimisation.",
        "A data protection impact assessment confirms lawful basis, minimisation and "
        "anonymisation for training data; the DPO signs off before training starts.",
        "Data Protection Officer",
        "UAE PDPL; EU AI Act Art. 10",
    ),
    (
        "AIRC-005",
        "Fairness & Bias",
        "Traditional",
        "Models produce discriminatory outcomes for protected groups (e.g. in recruitment).",
        "Bias testing across relevant groups is performed before release and quarterly in "
        "production against agreed fairness thresholds, with results reviewed by the model owner.",
        "Model Owner",
        "EU AI Act Art. 10; ISO/IEC 42001 Annex A",
    ),
    (
        "AIRC-006",
        "Human Oversight",
        "Both",
        "Automated outputs are acted on without meaningful human review.",
        "Decisions with legal or similarly significant effects require documented human review "
        "with authority to override, and reviewers are trained on automation bias.",
        "Business Process Owner",
        "EU AI Act Art. 14; UAE PDPL",
    ),
    (
        "AIRC-007",
        "GenAI Output Quality",
        "GenAI",
        "Generative models produce hallucinated or inaccurate content presented to customers.",
        "GenAI applications are grounded on approved knowledge sources (RAG), cite sources, and "
        "are evaluated against a groundedness test set before release and after each change.",
        "Product Owner",
        "ISO/IEC 42001 Annex A; NIST AI RMF MEASURE",
    ),
    (
        "AIRC-008",
        "GenAI Security",
        "GenAI",
        "Prompt injection or jailbreaks cause the model to leak data or take unintended actions.",
        "GenAI apps apply input/output filtering, treat retrieved content as untrusted data, "
        "restrict tool permissions to least privilege, and are red-teamed before launch.",
        "Information Security",
        "OWASP Top 10 for LLM Applications; ISO/IEC 42001 Annex A",
    ),
    (
        "AIRC-009",
        "Transparency",
        "GenAI",
        "Customers are not aware they are interacting with an AI system or AI-generated content.",
        "Customer-facing chatbots disclose that the user is talking to AI at the start of the "
        "conversation and offer a route to a human agent; synthetic content is labelled.",
        "Customer Experience",
        "EU AI Act Art. 50",
    ),
    (
        "AIRC-010",
        "Monitoring & Incident Management",
        "Both",
        "Model drift, failures or incidents go undetected in production.",
        "Production AI systems have monitoring for performance, drift and misuse, logs retained "
        "for at least six months, and incidents are reported through the AI incident process.",
        "Model Owner",
        "EU AI Act Art. 26; ISO/IEC 42001 Cl.9",
    ),
]


def make_docx(path: Path) -> None:
    d = docx.Document()
    d.core_properties.title = "ISO/IEC 42001 Overview (Sample)"
    d.add_heading("ISO/IEC 42001 Overview (Sample)", level=0)
    d.add_paragraph(
        "Placeholder paraphrase of ISO/IEC 42001:2023 (AI management systems) for testing Ben. "
        "Not the official text - use your licensed copy of the standard."
    )
    for heading, body in ISO_CLAUSES:
        d.add_heading(heading, level=1)
        d.add_paragraph(body)
    d.save(str(path))


def make_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, title="UAE PDPL and AI - Guidance Notes (Sample)"
    )
    story = [Paragraph("UAE PDPL and AI - Guidance Notes (Sample)", styles["Title"])]
    for i, page in enumerate(PDPL_PAGES):
        if i:
            story.append(PageBreak())
        for heading, body in page:
            story += [Paragraph(heading, styles["Heading2"]), Paragraph(body, styles["BodyText"])]
            story.append(Spacer(1, 12))
    doc.build(story)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Controls"
    header = [
        "control_id",
        "risk_category",
        "ai_type",
        "risk_description",
        "control_description",
        "owner",
        "framework_mapping",
    ]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in CONTROLS:
        ws.append(list(row))
    wb.save(str(path))


if __name__ == "__main__":
    KNOWLEDGE.mkdir(exist_ok=True)
    make_docx(KNOWLEDGE / "sample_iso42001_overview.docx")
    make_pdf(KNOWLEDGE / "sample_uae_pdpl_notes.pdf")
    make_xlsx(KNOWLEDGE / "ai_risk_control_library.xlsx")
    print("Sample knowledge written to", KNOWLEDGE)
